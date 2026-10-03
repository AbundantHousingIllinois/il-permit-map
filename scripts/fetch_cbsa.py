#!/usr/bin/env python3
"""Fetch the OMB metropolitan and micropolitan area delineation (ROADMAP item 2).

  www2.census.gov/programs-surveys/metro-micro/geographies/reference-files/
      2023/delineation-files/list1_2023.xlsx
      -- OMB Bulletin 23-01 (July 2023), as published by the Census Bureau:
         every core-based statistical area (CBSA), metropolitan or
         micropolitan, and the counties that make it up.

The Metros view uses it to group Illinois counties into areas, counting only
the Illinois counties of an area that crosses a state line (Austin, 2026-10-02).
BPS records carry a CBSA code of their own, but it follows whichever delineation
was in force that year; one fixed delineation keeps every year comparable.

Idempotent: an already-downloaded file is left alone.
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import bps_layout as L

URL = ("https://www2.census.gov/programs-surveys/metro-micro/geographies/"
       "reference-files/2023/delineation-files/list1_2023.xlsx")


def main() -> int:
    print("=" * 78)
    print("Fetching the OMB metro/micro area delineation (July 2023)")
    print("=" * 78)
    L.download(URL, L.CBSA_DELINEATION,
               "OMB Bulletin 23-01 CBSA delineation (counties per metro/micro area)")

    from openpyxl import load_workbook
    ws = load_workbook(L.CBSA_DELINEATION, read_only=True).worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    head = next(i for i, r in enumerate(rows) if r and "CBSA Code" in [str(c) for c in r])
    cols = [str(c) for c in rows[head]]
    st, cbsa, kind = cols.index("FIPS State Code"), cols.index("CBSA Code"), \
        cols.index("Metropolitan/Micropolitan Statistical Area")
    il = [r for r in rows[head + 1:] if r and str(r[st]).zfill(2) == L.STATE_FIPS]
    areas = {r[cbsa]: r[kind] for r in il}
    print(f"\n  Illinois counties in a CBSA : {len(il)}")
    print(f"  CBSAs with an Illinois county: {len(areas)} "
          f"({sum('Metro' in v for v in areas.values())} metropolitan, "
          f"{sum('Micro' in v for v in areas.values())} micropolitan)")
    print(f"Provenance appended to {L.rel(L.SOURCES_MD)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
